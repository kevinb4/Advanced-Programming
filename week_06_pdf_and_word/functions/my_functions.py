import docx
import PyPDF4
import openpyxl
from bs4 import BeautifulSoup

def load_excel_data(path):
    """Load the excel data from the specified path
    Arguments:
        path [string] -- the path to the excel file
    Returns:
        [list] -- the data from the excel file"""
    wb = openpyxl.load_workbook(path)
    sheet = wb.active

    # loop through all rows and columns to get the data
    data = []
    for sheet_row in sheet.iter_rows():
        row = []
        for cell in sheet_row:
            row.append(str(cell.value).replace("\u202c", "")) # remove the encoding characters

        data.append(row)

    return data

def load_html_data(path):
    """Load the html data from the specified path
    Arguments:
        path [string] -- the path to the html file
    Returns:
        [string] -- the data from the html file"""
    html = open(path, encoding="utf8")
    it_soup = BeautifulSoup(html, "lxml")
    paragraphs = it_soup.find_all("p")
    data = ""

    for paragraph in paragraphs: # loop through all paragraphs to extract the text cleaned up
        data += f"{paragraph.getText()}\n\n"

    return data[:-2] # get rid of last two new lines

def load_pdf_data(path):
    """Load the pdf data from the specified path
    Arguments:
        path [string] -- the path to the pdf file
    Returns:
        [string] -- the data from the pdf file"""
    pdf = open(path, "rb")
    pdf_reader = PyPDF4.PdfFileReader(pdf)

    # split the text by the period and new line
    split_text = pdf_reader.getPage(0).extractText().strip().split(".\n \n")
    patched_text = ""

    for text in range(len(split_text)): # loop through the split array to disect what we need and shove it back in
        item = split_text.pop(0)

        if "!\n \n" in item: # find the ! since one of the paragraphs ends with it
            lines = item.split("!\n \n")
            length = len(lines)
            for line in range(length): # loop through the new split
                final = ""
                # give all but the last line a ! since it was split on that
                if line == length - 1:
                    lines[line] += "."
                else:
                    lines[line] += "!"

            final = ""
            for line in lines: # loop through everything to remove the random line breaks and add two for the new paragraph
                final += line.replace("\n", "") + "\n\n"
            patched_text += final

            # could also check for ? as well, but since the current doc doesn't have it, I will leave it as is
        else:
            # the original split just needs to have the random line breaks removed and a period added with two new lines
            patched_text += item.replace("\n", "") + ".\n\n"

    # remove the last two new lines and the extra period
    return patched_text[:-3]

def load_word_data(path):
    """Load the word data from the specified path
    Arguments:
        path [string] -- the path to the word file
    Returns:
        [string] -- the data from the word file"""
    doc = docx.Document(path)
    full_text = ""

    for para in doc.paragraphs: # loop through all paragraphs
        full_text += f"{para.text}\n\n"

    return full_text[:-2] # get rid of trailing new lines

def write_table_page(doc, title, data, last=False):
    """Write the page title and table data
    Arguments:
        doc [object] -- the word document to write to
        title [string] -- the title of the page
        data [list] -- the data to insert into the table
        last [boolean] -- whether or not this is the last page to write (so a page break is not added)"""
    doc.add_heading(title, 0)

    # create a table the size of the imported data
    table = doc.add_table(rows=len(data), cols=len(data[0]))

    # loop through imported data and insert into table by first the row and then the column
    for row in range(len(data)):
        for column in range(len(data[row])):
            table.cell(row, column).text = data[row][column]

    if not last:
        doc.add_page_break()

def write_text_page(doc, title, data, last=False):
    """Write the page title and data
    Arguments:
        doc [object] -- the word document to write to
        title [string] -- the title of the page to write
        data [string] -- the data to insert into the page
        last [boolean] -- whether or not this is the last page to write (so a page break is not added)"""
    doc.add_heading(title, 0)
    doc.add_paragraph(data)

    if not last:
        doc.add_page_break()